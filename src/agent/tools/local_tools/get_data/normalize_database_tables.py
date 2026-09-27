from __future__ import annotations

import argparse
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.project_paths import DATA_ROOT

PROCESSED_DIR = DATA_ROOT / "processed"

TREE_COLUMNS = [
    "tree_id",
    "species",
    "plot_id",
    "tree_position_x_m",
    "tree_position_y_m",
    "tree_position_z_m",
    "elevation_m",
]

MEASUREMENT_COLUMNS = [
    "tree_id",
    "measurement_year",
    "dbh_m",
    "tree_height_m",
    "crown_diameter_m",
    "crown_diameter_ns_m",
    "crown_diameter_ew_m",
    "crown_area_m2",
    "crown_volume_m3",
    "crown_base_height_m",
    "data_source",
]

MEASUREMENT_FIELDS_2025 = {
    "tree_height_m_2025": "tree_height_m",
    "crown_diameter_m_2025": "crown_diameter_m",
    "crown_diameter_ns_m_2025": "crown_diameter_ns_m",
    "crown_diameter_ew_m_2025": "crown_diameter_ew_m",
    "crown_area_m2_2025": "crown_area_m2",
    "crown_volume_m3_2025": "crown_volume_m3",
    "crown_base_height_m_2025": "crown_base_height_m",
}


def default_source() -> Path:
    candidates = sorted(
        path
        for path in PROCESSED_DIR.glob("*_EN.xlsx")
        if not path.name.startswith("~$") and not path.stem.endswith("_DB")
    )
    if len(candidates) != 1:
        raise FileNotFoundError(
            "Expected exactly one source workbook matching *_EN.xlsx; "
            f"found {len(candidates)}. Pass --source explicitly."
        )
    return candidates[0]


def build_tables(source: Path) -> tuple[pd.DataFrame, pd.DataFrame]:
    flat = pd.read_excel(
        source,
        sheet_name="matched_trees",
        dtype={"tree_id": "string", "plot_id": "string"},
    )
    required = set(TREE_COLUMNS) | {
        "dbh_m_2016",
        "dbh_m_2021",
        "dbh_m_2025",
        *MEASUREMENT_FIELDS_2025,
    }
    missing = sorted(required.difference(flat.columns))
    if missing:
        raise ValueError(f"Source workbook is missing columns: {missing}")

    trees = flat[TREE_COLUMNS].copy()
    trees["tree_id"] = trees["tree_id"].str.strip()

    frames: list[pd.DataFrame] = []
    for year, source_name in (
        (2016, "field_inventory"),
        (2021, "field_inventory"),
        (2025, "ue_segmentation"),
    ):
        measurements = pd.DataFrame(index=flat.index)
        measurements["tree_id"] = flat["tree_id"].str.strip()
        measurements["measurement_year"] = year
        measurements["dbh_m"] = flat[f"dbh_m_{year}"]
        for target in MEASUREMENT_FIELDS_2025.values():
            measurements[target] = float("nan")
        if year == 2025:
            for source_column, target_column in MEASUREMENT_FIELDS_2025.items():
                measurements[target_column] = flat[source_column]
        measurements["data_source"] = source_name
        frames.append(measurements[MEASUREMENT_COLUMNS])

    tree_measurements = pd.concat(frames, ignore_index=True)
    tree_measurements = tree_measurements.sort_values(
        ["tree_id", "measurement_year"], kind="stable"
    ).reset_index(drop=True)

    if trees["tree_id"].isna().any() or trees["tree_id"].eq("").any():
        raise ValueError("trees contains an empty tree_id")
    if trees["tree_id"].duplicated().any():
        raise ValueError("trees contains duplicate tree_id values")
    if tree_measurements.duplicated(["tree_id", "measurement_year"]).any():
        raise ValueError("tree_measurements contains duplicate (tree_id, measurement_year)")
    if not set(tree_measurements["tree_id"]).issubset(set(trees["tree_id"])):
        raise ValueError("tree_measurements contains a tree_id not present in trees")
    if tree_measurements["dbh_m"].isna().any():
        raise ValueError("tree_measurements contains an empty dbh_m")

    return trees, tree_measurements


def style_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center")
        for column_index, cells in enumerate(sheet.columns, start=1):
            values = [str(cell.value) if cell.value is not None else "" for cell in cells[:200]]
            width = min(max(max(map(len, values), default=8) + 2, 10), 30)
            sheet.column_dimensions[get_column_letter(column_index)].width = width
    workbook.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Split the flat English tree workbook into normalized database tables."
    )
    parser.add_argument("--source", type=Path, default=None)
    parser.add_argument("--output", type=Path, default=None)
    args = parser.parse_args()

    source = (args.source or default_source()).resolve()
    output = (args.output or source.with_name(f"{source.stem}_DB.xlsx")).resolve()
    trees_csv = output.with_name(f"{output.stem}_trees.csv")
    measurements_csv = output.with_name(f"{output.stem}_tree_measurements.csv")

    trees, tree_measurements = build_tables(source)
    output.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        trees.to_excel(writer, sheet_name="trees", index=False)
        tree_measurements.to_excel(writer, sheet_name="tree_measurements", index=False)
    style_workbook(output)

    # UTF-8 with BOM opens cleanly in Excel and is accepted by PostgreSQL/pgAdmin.
    trees.to_csv(trees_csv, index=False, encoding="utf-8-sig")
    tree_measurements.to_csv(measurements_csv, index=False, encoding="utf-8-sig")

    print(f"source={source}")
    print(f"workbook={output}")
    print(f"trees_csv={trees_csv}")
    print(f"tree_measurements_csv={measurements_csv}")
    print(f"trees={len(trees)}")
    print(f"tree_measurements={len(tree_measurements)}")


if __name__ == "__main__":
    main()
