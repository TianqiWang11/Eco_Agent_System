from __future__ import annotations
import argparse
import warnings
from collections import defaultdict, deque
from pathlib import Path
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from scipy.optimize import linear_sum_assignment
from scipy.spatial import cKDTree  # pyright: ignore[reportAttributeAccessIssue]

from src.project_paths import DATA_ROOT, REFERENCE_DATA_DIR

CORR = REFERENCE_DATA_DIR / "副本建模每木与原始数据对应表.xlsx"
MONITOR = REFERENCE_DATA_DIR / "车八岭20公顷样地监测数据（2016-2022）-final.xlsx"
SEGMENT = REFERENCE_DATA_DIR / "20公顷大样地单木分割属性合并.xlsx"
OUTPUT = DATA_ROOT / "processed" / "车八岭20公顷样地单木对应表_2016_2021_2025_EN.xlsx"

RESULT_FIELDS = {
    "编号": "tree_id",
    "树种": "species",
    "样地编号": "plot_id",
    "树方位X(米)": "tree_position_x_m",
    "树方位Y(米)": "tree_position_y_m",
    "树方位Z(米)": "tree_position_z_m",
    "树高(米)": "tree_height_m_2025",
    "胸径2016(米)": "dbh_m_2016",
    "胸径2021(米)": "dbh_m_2021",
    "胸径2025(米)": "dbh_m_2025",
    "冠径(米)": "crown_diameter_m_2025",
    "南北冠径(米)": "crown_diameter_ns_m_2025",
    "东西冠径(米)": "crown_diameter_ew_m_2025",
    "冠幅面积(平方米)": "crown_area_m2_2025",
    "冠幅体积(立方米)": "crown_volume_m3_2025",
    "枝下高(米)": "crown_base_height_m_2025",
    "海拔(米)": "elevation_m",
}

AUDIT_FIELDS = {
    "编号": "tree_id",
    "树种": "species",
    "监测tag": "monitoring_tag",
    "监测树种": "monitoring_species",
    "树种是否一致": "species_match",
    "转换后东坐标": "transformed_easting",
    "转换后北坐标": "transformed_northing",
    "样地编号": "plot_id",
    "树方位X(米)": "tree_position_x_m",
    "树方位Y(米)": "tree_position_y_m",
    "坐标匹配距离(米)": "match_distance_m",
    "对应表源行": "correspondence_source_row",
    "单木分割表源行": "segmentation_source_row",
}

EXCLUSION_FIELDS = {
    "编号": "tree_id",
    "树种": "species",
    "转换后东坐标": "transformed_easting",
    "转换后北坐标": "transformed_northing",
    "未纳入原因": "exclusion_reason",
    "对应表源行": "correspondence_source_row",
}


def text_id(value):
    if pd.isna(value):
        return ""
    value = str(value).strip()
    return value[:-2] if value.endswith(".0") else value


def match_key(value):
    return text_id(value).lstrip("0") or "0"


def sources():
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Cannot parse header or footer.*")
        raw = pd.read_excel(CORR, sheet_name="建模每木", header=None)
        monitor = pd.read_excel(MONITOR, sheet_name=0)
        segment = pd.read_excel(SEGMENT, sheet_name="合并数据")
    corr = pd.DataFrame({
        "_source_row": np.arange(3, len(raw) + 1),
        "编号": raw.iloc[2:, 0].map(text_id).to_numpy(),
        "树种": raw.iloc[2:, 1].astype("string").str.strip().to_numpy(),
        "_north": pd.to_numeric(raw.iloc[2:, 8], errors="coerce").to_numpy(),
        "_east": pd.to_numeric(raw.iloc[2:, 9], errors="coerce").to_numpy(),
    })
    corr = corr[corr["编号"].ne("")].copy()
    corr["_key"] = corr["编号"].map(match_key)
    return corr, monitor, segment


def trunks(monitor):
    trunk = monitor[pd.to_numeric(monitor["branch"], errors="coerce").eq(0)].copy()
    trunk["_key"] = trunk["tag"].map(match_key)
    trunk["_dbh16"] = pd.to_numeric(trunk["dbh(2016)"], errors="coerce")
    trunk["_dbh21"] = pd.to_numeric(trunk["dbh1(2021)"], errors="coerce")
    trunk = trunk[~trunk["_key"].isin({"Null", "null", "nan"})].copy()
    all_keys = set(trunk["_key"])
    valid = trunk[trunk["_dbh16"].notna() & trunk["_dbh21"].notna()].copy()
    dup = valid["_key"].duplicated(False)
    if dup.any():
        distinct = (valid.loc[dup, ["_key", "spname", "_dbh16", "_dbh21"]]
                    .drop_duplicates().groupby("_key").size())
        if (distinct > 1).any():
            raise ValueError("监测表存在同一tag的冲突branch0记录")
        valid = valid.drop_duplicates("_key")
    return valid, all_keys


def candidates(corr, monitor):
    cols = ["_key", "tag", "spname", "_dbh16", "_dbh21"]
    joined = corr.merge(monitor[cols], on="_key", how="inner").rename(
        columns={"tag": "监测tag", "spname": "监测树种"})
    joined["_species_same"] = (
        joined["树种"].fillna("").astype(str).str.strip()
        == joined["监测树种"].fillna("").astype(str).str.strip())
    joined = joined.sort_values(
        ["_key", "_species_same", "_source_row"], ascending=[True, False, True])
    unique = joined.drop_duplicates("_key").dropna(subset=["_east", "_north"])
    return unique.reset_index(drop=True), set(unique["_source_row"].astype(int))


def coordinate_assignment(trees, segment, radius):
    sx = pd.to_numeric(segment["树方位X(米)"], errors="coerce")
    sy = pd.to_numeric(segment["树方位Y(米)"], errors="coerce")
    valid = segment.loc[sx.notna() & sy.notna()].copy().reset_index()
    coords = np.column_stack([pd.to_numeric(valid["树方位X(米)"]),
                              pd.to_numeric(valid["树方位Y(米)"])])
    target = np.column_stack([trees["_east"], trees["_north"]])
    possible = cKDTree(coords).query_ball_point(target, r=radius)
    reverse = defaultdict(list)
    for left, rights in enumerate(possible):
        for right in rights:
            reverse[int(right)].append(left)
    matches, visited = {}, set()
    possible_left = {left for left, rights in enumerate(possible) if rights}
    for seed in sorted(possible_left):
        if seed in visited:
            continue
        lefts, rights, queue = set(), set(), deque([("L", seed)])
        while queue:
            side, node = queue.popleft()
            if side == "L":
                if node in lefts:
                    continue
                lefts.add(node); visited.add(node)
                queue.extend(("R", int(r)) for r in possible[node])
            else:
                if node in rights:
                    continue
                rights.add(node)
                queue.extend(("L", int(l)) for l in reverse[node])
        lefts, rights = sorted(lefts), sorted(rights)
        rpos = {value: pos for pos, value in enumerate(rights)}
        cost = np.full((len(lefts), len(rights) + len(lefts)), 1_000_000.0)
        cost[:, len(rights):] = radius + 1
        for row, left in enumerate(lefts):
            for right in possible[left]:
                cost[row, rpos[int(right)]] = np.linalg.norm(target[left] - coords[int(right)])
        rows, columns = linear_sum_assignment(cost)
        for row, column in zip(rows, columns):
            if column < len(rights) and cost[row, column] <= radius:
                left, right = lefts[int(row)], rights[int(column)]
                matches[left] = (int(valid.iloc[right]["index"]), float(cost[row, column]))
    return matches, possible_left


def output_frames(trees, segment, matches):
    result, audit = [], []
    for left in sorted(matches):
        seg_index, distance = matches[left]
        tree, seg = trees.iloc[left], segment.loc[seg_index]
        row = {"编号": tree["编号"], "树种": tree["树种"]}
        for column in segment.columns:
            if column == "树编号":
                continue
            if column == "胸径(米)":
                row["胸径2016(米)"] = float(tree["_dbh16"]) / 100
                row["胸径2021(米)"] = float(tree["_dbh21"]) / 100
                row["胸径2025(米)"] = pd.to_numeric(seg[column], errors="coerce")
            else:
                row[column] = seg[column]
        result.append(row)
        audit.append({
            "编号": tree["编号"], "树种": tree["树种"],
            "监测tag": tree["监测tag"], "监测树种": tree["监测树种"],
            "树种是否一致": bool(tree["_species_same"]),
            "转换后东坐标": tree["_east"], "转换后北坐标": tree["_north"],
            "样地编号": seg["样地编号"],
            "树方位X(米)": seg["树方位X(米)"], "树方位Y(米)": seg["树方位Y(米)"],
            "坐标匹配距离(米)": distance, "对应表源行": int(tree["_source_row"]),
            "单木分割表源行": seg_index + 2,
        })
    return pd.DataFrame(result), pd.DataFrame(audit)


def exclusions(corr, all_keys, valid_keys, selected, eligible, matches, possible):
    positions = {int(row["_source_row"]): pos for pos, (_, row) in enumerate(eligible.iterrows())}
    rows = []
    for _, row in corr.iterrows():
        source_row, tree_key = int(row["_source_row"]), str(row["_key"])
        if tree_key not in all_keys:
            reason = "完整编号没有对应的branch0监测记录"
        elif tree_key not in valid_keys:
            reason = "branch0的2016或2021年DBH为空"
        elif source_row not in selected:
            reason = "对应表编号重复，已保留物种一致或更早的一条"
        else:
            pos = positions[source_row]
            if pos in matches:
                continue
            reason = ("2米内存在候选，但一对一分配冲突"
                      if pos in possible else "最近单木分割坐标超过2米")
        rows.append({"编号": row["编号"], "树种": row["树种"],
                     "转换后东坐标": row["_east"], "转换后北坐标": row["_north"],
                     "未纳入原因": reason, "对应表源行": source_row})
    return pd.DataFrame(rows)


def style(path):
    book = load_workbook(path)
    fill = PatternFill("solid", fgColor="1F4E78")
    for sheet in book.worksheets:
        sheet.freeze_panes = "A2"; sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = fill; cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center")
        for column_index, cells in enumerate(sheet.columns, start=1):
            values = [str(c.value) if c.value is not None else "" for c in cells[:200]]
            sheet.column_dimensions[get_column_letter(column_index)].width = min(
                max(max(map(len, values), default=8) + 2, 10), 30)
    book.save(path)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=OUTPUT)
    parser.add_argument("--radius", type=float, default=2.0)
    args = parser.parse_args()
    corr, monitor, segment = sources()
    valid_monitor, all_keys = trunks(monitor)
    eligible, selected = candidates(corr, valid_monitor)
    matches, possible = coordinate_assignment(eligible, segment, args.radius)
    result, audit = output_frames(eligible, segment, matches)
    excluded = exclusions(corr, all_keys, set(valid_monitor["_key"]),
                          selected, eligible, matches, possible)
    dbh = ["胸径2016(米)", "胸径2021(米)", "胸径2025(米)"]
    segment_rows = audit["单木分割表源行"]
    if (result["编号"].duplicated().any() or segment_rows.duplicated().any()
            or result[dbh].isna().any().any()):
        raise ValueError("结果存在重复树或空DBH")
    summary = pd.DataFrame([
        ("对应表原始记录", len(corr)),
        ("branch0且2016/2021 DBH均有效的监测树", len(valid_monitor)),
        ("完整编号连接并去重后的候选树", len(eligible)),
        (f"{args.radius:g}米内一对一坐标匹配成功", len(result)),
        ("未纳入记录", len(excluded)),
        ("成功记录中树种名称一致", int(audit["树种是否一致"].sum())),
        ("最大坐标匹配距离(米)", float(audit["坐标匹配距离(米)"].max())),
    ], columns=["指标", "数值"])
    rules = pd.DataFrame([
        ("编号", "对应表编号去前导零后匹配完整tag；不盲取后6位，避免跨样方重号"),
        ("主干与空值", "只取branch=0；2016或2021 DBH任一为空即剔除"),
        ("坐标", f"东/北坐标对应树方位X/Y；{args.radius:g}米内全局一对一最小距离分配"),
        ("DBH单位", "三年均为米；2016/2021原始厘米除以100，2025保留分割表原值"),
    ], columns=["规则", "说明"])
    args.output.parent.mkdir(parents=True, exist_ok=True)
    result = result.rename(columns=RESULT_FIELDS)
    audit = audit.rename(columns=AUDIT_FIELDS)
    excluded = excluded.rename(columns=EXCLUSION_FIELDS)
    summary = summary.rename(columns={"指标": "metric", "数值": "value"})
    rules = rules.rename(columns={"规则": "rule", "说明": "description"})
    with pd.ExcelWriter(args.output, engine="openpyxl") as writer:
        result.to_excel(writer, sheet_name="matched_trees", index=False)
        audit.to_excel(writer, sheet_name="match_audit", index=False)
        excluded.to_excel(writer, sheet_name="excluded_records", index=False)
        summary.to_excel(writer, sheet_name="notes", index=False)
        rules.to_excel(writer, sheet_name="notes", index=False, startrow=len(summary) + 3)
    style(args.output)
    print(f"output={args.output}")
    print(f"matched={len(result)} excluded={len(excluded)}")
    print(f"max_distance={audit['match_distance_m'].max():.6f}")


if __name__ == "__main__":
    main()
