"""Species-aware growth prediction and in-memory Excel artifact generation."""

from __future__ import annotations

import base64
import os
from io import BytesIO
from typing import Any

import numpy as np
import pandas as pd
import psycopg
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from scipy.stats import linregress

from .model_store import load_bundle
from .training import FEATURE_COLUMNS, add_competition_features


EXCEL_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
SHAPE_COLUMNS = [
    "tree_height_m",
    "crown_diameter_m",
    "crown_diameter_ns_m",
    "crown_diameter_ew_m",
    "crown_area_m2",
    "crown_volume_m3",
]
OUTPUT_COLUMN_WIDTHS = [
    23,
    29,
    28.36328125,
    21.453125,
    38.54296875,
    37.08984375,
    44.36328125,
    53.6328125,
    46.54296875,
    43,
    39.26953125,
]


def output_headers(target_year: int) -> list[str]:
    """The canonical prediction workbook template, stored in code."""
    return [
        "编号",
        "物种名称",
        "预测年份",
        "样地编号",
        f"predicted_dbh_m_{target_year}",
        f"estimated_tree_height_m_{target_year}",
        f"estimated_crown_diameter_m_{target_year}",
        f"estimated_crown_diameter_ns_m_{target_year}",
        f"estimated_crown_diameter_ew_m_{target_year}",
        f"estimated_crown_area_m2_{target_year}",
        f"estimated_crown_volume_m3_{target_year}",
    ]


PREDICT_SCHEMA = {
    "type": "object",
    "properties": {
        "operation": {
            "type": "string",
            "enum": ["status", "predict_tree", "predict_species", "predict_species_trees", "predict_sample"],
        },
        "tree_id": {"type": "string", "minLength": 1, "maxLength": 64},
        "species": {"type": "string", "minLength": 1, "maxLength": 100},
        "dbh_m": {"type": "number", "exclusiveMinimum": 0, "maximum": 5},
        "base_year": {"type": "integer", "enum": [2016, 2021]},
        "target_year": {"type": "integer", "minimum": 2026, "maximum": 2100},
        "sample_size": {"type": "integer", "minimum": 1, "maximum": 200},
        "seed": {"type": "integer", "minimum": 0, "maximum": 2147483647},
        "elevation_m": {"type": "number", "minimum": -500, "maximum": 9000},
        "tree_position_x_m": {"type": "number"},
        "tree_position_y_m": {"type": "number"},
        "neighbor_count_10m": {"type": "number", "minimum": 0},
        "neighbor_basal_area_m2_10m": {"type": "number", "minimum": 0},
        "hegyi_index_10m": {"type": "number", "minimum": 0},
    },
    "required": ["operation"],
    "additionalProperties": False,
}


def _database_frame() -> pd.DataFrame:
    url = os.getenv("DATABASE_URL", "").strip()
    if not url:
        raise RuntimeError("DATABASE_URL is not configured")
    query = """
        SELECT t.tree_id, t.species, t.plot_id, t.tree_position_x_m,
               t.tree_position_y_m, t.elevation_m,
               max(m.dbh_m) FILTER (WHERE m.measurement_year=2016) AS dbh_m_2016,
               max(m.dbh_m) FILTER (WHERE m.measurement_year=2021) AS dbh_m_2021,
               max(m.dbh_m) FILTER (WHERE m.measurement_year=2025) AS dbh_m_2025,
               max(m.tree_height_m) FILTER (WHERE m.measurement_year=2025) AS tree_height_m,
               max(m.crown_diameter_m) FILTER (WHERE m.measurement_year=2025) AS crown_diameter_m,
               max(m.crown_diameter_ns_m) FILTER (WHERE m.measurement_year=2025) AS crown_diameter_ns_m,
               max(m.crown_diameter_ew_m) FILTER (WHERE m.measurement_year=2025) AS crown_diameter_ew_m,
               max(m.crown_area_m2) FILTER (WHERE m.measurement_year=2025) AS crown_area_m2,
               max(m.crown_volume_m3) FILTER (WHERE m.measurement_year=2025) AS crown_volume_m3
        FROM public.trees t JOIN public.tree_measurements m USING(tree_id)
        GROUP BY t.tree_id, t.species, t.plot_id, t.tree_position_x_m,
                 t.tree_position_y_m, t.elevation_m
        ORDER BY t.tree_id
    """
    with psycopg.connect(url) as connection:
        cursor = connection.execute(query)
        rows = cursor.fetchall()
        description = cursor.description
        if description is None:
            raise RuntimeError("PostgreSQL did not return tree columns")
        columns = [column.name for column in description]
    return pd.DataFrame(rows, columns=columns)


def _prepared_database(base_year: int) -> pd.DataFrame:
    frame = _database_frame()
    dbh_column = f"dbh_m_{base_year}"
    frame = frame.dropna(subset=["species", "plot_id", "dbh_m_2016", "dbh_m_2021", dbh_column])
    frame = frame.loc[(frame["dbh_m_2016"] > 0) & (frame["dbh_m_2021"] > 0) & (frame[dbh_column] > 0)]
    if frame.empty:
        raise ValueError("数据库中没有可用于预测的树木")
    return add_competition_features(frame.reset_index(drop=True), dbh_column)


def _shape_relations(frame: pd.DataFrame) -> dict[str, dict[str, float]]:
    relations = {}
    for column in SHAPE_COLUMNS[:4]:
        eligible = frame.loc[(frame["dbh_m_2025"] > 0) & (frame[column] > 0)]
        if len(eligible) < 100:
            raise ValueError(f"2025年{column}数据不足，无法估算形态")
        fitted: Any = linregress(
            np.log(eligible["dbh_m_2025"]), np.log(eligible[column])
        )
        exponent = fitted.slope
        intercept = fitted.intercept
        rvalue = fitted.rvalue
        if not np.isfinite(exponent) or exponent <= 0:
            raise ValueError(f"{column}异速生长关系不可用")
        relations[column] = {
            "intercept": float(intercept),
            "exponent": float(exponent),
            "r2": float(rvalue**2),
        }
    return relations


def _prediction_payload(feature_row: pd.DataFrame, base_year: int, target_years: list[int]) -> dict:
    model, metadata = load_bundle()
    for column in FEATURE_COLUMNS:
        if column == "species":
            feature_row[column] = feature_row[column].fillna("unknown").astype(str)
        else:
            median = metadata["numeric_feature_medians"][column]
            feature_row[column] = pd.to_numeric(feature_row[column], errors="coerce").fillna(median)
    annual_growth = float(model.predict(feature_row[FEATURE_COLUMNS])[0])
    annual_growth = float(np.clip(annual_growth, 0.0, metadata["target_upper_cap_m_per_year"]))
    base_dbh = float(feature_row.iloc[0]["dbh_m"])
    predictions = []
    for year in sorted(target_years):
        if year <= base_year:
            raise ValueError("target_year必须晚于base_year")
        predictions.append({
            "target_year": year,
            "predicted_dbh_m": base_dbh + annual_growth * (year - base_year),
        })
    species = str(feature_row.iloc[0]["species"])
    return {
        "model_name": metadata["model_name"],
        "model_version": metadata["model_version"],
        "base_year": base_year,
        "base_dbh_m": base_dbh,
        "annual_dbh_growth_m": annual_growth,
        "species": species,
        "species_seen_in_training": species in metadata["known_species"],
        "predictions": predictions,
        "experimental": True,
    }


def _shape_estimates(
    tree: pd.Series,
    predicted_dbh_2025: float,
    predicted_dbh_target: float,
    relations: dict[str, dict[str, float]],
) -> dict[str, float]:
    observed = all(pd.notna(tree.get(column)) and float(tree[column]) > 0 for column in SHAPE_COLUMNS)
    if observed:
        ratio = predicted_dbh_target / predicted_dbh_2025
        factors = {
            column: ratio ** relations[column]["exponent"]
            for column in SHAPE_COLUMNS[:4]
        }
        area_factor = factors["crown_diameter_ns_m"] * factors["crown_diameter_ew_m"]
        volume_factor = area_factor * factors["tree_height_m"]
        return {
            "tree_height_m": float(tree["tree_height_m"]) * factors["tree_height_m"],
            "crown_diameter_m": float(tree["crown_diameter_m"]) * factors["crown_diameter_m"],
            "crown_diameter_ns_m": float(tree["crown_diameter_ns_m"]) * factors["crown_diameter_ns_m"],
            "crown_diameter_ew_m": float(tree["crown_diameter_ew_m"]) * factors["crown_diameter_ew_m"],
            "crown_area_m2": float(tree["crown_area_m2"]) * area_factor,
            "crown_volume_m3": float(tree["crown_volume_m3"]) * volume_factor,
        }

    estimates = {
        column: float(np.exp(values["intercept"]) * predicted_dbh_target ** values["exponent"])
        for column, values in relations.items()
    }
    estimates["crown_area_m2"] = (
        np.pi * estimates["crown_diameter_ns_m"] * estimates["crown_diameter_ew_m"] / 4.0
    )
    estimates["crown_volume_m3"] = estimates["crown_area_m2"] * estimates["tree_height_m"] * 0.5
    return estimates


def _output_record(
    tree: pd.Series,
    result: dict,
    target_year: int,
    relations: dict[str, dict[str, float]],
) -> dict:
    by_year = {item["target_year"]: item["predicted_dbh_m"] for item in result["predictions"]}
    shape = _shape_estimates(tree, by_year[2025], by_year[target_year], relations)
    headers = output_headers(target_year)
    return dict(zip(headers, [
        str(tree.get("tree_id") or "scenario-001"),
        str(tree["species"]),
        target_year,
        str(tree.get("plot_id") or ""),
        by_year[target_year],
        shape["tree_height_m"],
        shape["crown_diameter_m"],
        shape["crown_diameter_ns_m"],
        shape["crown_diameter_ew_m"],
        shape["crown_area_m2"],
        shape["crown_volume_m3"],
    ]))


def _excel_artifact(records: list[dict], target_year: int, filename: str) -> dict:
    headers = output_headers(target_year)
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = f"predictions_{target_year}"
    thin = Side(style="thin", color="000000")
    for column, (header, width) in enumerate(zip(headers, OUTPUT_COLUMN_WIDTHS), start=1):
        cell = sheet.cell(row=1, column=column, value=header)
        cell.font = Font(name="宋体", size=11, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="top")
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        sheet.column_dimensions[cell.column_letter].width = width
    for row, record in enumerate(records, start=2):
        for column, header in enumerate(headers, start=1):
            cell = sheet.cell(row=row, column=column, value=record[header])
            cell.font = Font(name="宋体", size=11)
            if column == 1:
                cell.number_format = "@"
            elif column >= 5:
                cell.number_format = "0.0000"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:K{len(records) + 1}"
    buffer = BytesIO()
    workbook.save(buffer)
    return {
        "title": f"{target_year}年树木预测数据",
        "mime": EXCEL_MIME,
        "filename": filename,
        "content_base64": base64.b64encode(buffer.getvalue()).decode("ascii"),
    }


def _artifact_result(operation: str, records: list[dict], result: dict, target_year: int, **extra) -> dict:
    filename = (
        f"dbh_predictions_{target_year}_random_{len(records)}.xlsx"
        if operation == "predict_sample"
        else f"dbh_predictions_{target_year}_{operation}.xlsx"
    )
    ue_actions = []
    if operation != "predict_species":
        for record in records:
            tree_id = str(record["编号"])
            if not tree_id or tree_id == "scenario-001":
                continue
            ue_actions.append({
                "type": "apply_prediction",
                "target_id": tree_id,
                "species": str(record["物种名称"]),
                "year": target_year,
                "dbh_m": float(record[f"predicted_dbh_m_{target_year}"]),
                "tree_height_m": float(record[f"estimated_tree_height_m_{target_year}"]),
                "crown_diameter_ns_m": float(record[f"estimated_crown_diameter_ns_m_{target_year}"]),
                "crown_diameter_ew_m": float(record[f"estimated_crown_diameter_ew_m_{target_year}"]),
                "crown_volume_m3": float(record[f"estimated_crown_volume_m3_{target_year}"]),
                "focus": False,
            })
        if ue_actions:
            ue_actions[0]["focus"] = True
    return {
        "operation": operation,
        "target_year": target_year,
        "row_count": len(records),
        "model_version": result["model_version"],
        "experimental": True,
        "preview": records[:5],
        "artifacts": [_excel_artifact(records, target_year, filename)],
        "ue_actions": ue_actions,
        **extra,
    }


def _predict_tree(arguments: dict) -> dict:
    tree_id = str(arguments.get("tree_id") or "").strip()
    target_year = arguments.get("target_year")
    if not tree_id or target_year is None:
        raise ValueError("predict_tree必须提供tree_id和target_year")
    base_year = int(arguments.get("base_year", 2021))
    frame = _prepared_database(base_year)
    row = frame.loc[frame["tree_id"].astype(str).eq(tree_id)]
    if row.empty:
        raise ValueError("数据库中没有找到该tree_id")
    relations = _shape_relations(frame)
    result = _prediction_payload(row[FEATURE_COLUMNS].copy(), base_year, [2025, int(target_year)])
    record = _output_record(row.iloc[0], result, int(target_year), relations)
    return _artifact_result("predict_tree", [record], result, int(target_year), tree_id=tree_id)


def _predict_species(arguments: dict) -> dict:
    species = str(arguments.get("species") or "").strip()
    target_year = arguments.get("target_year")
    if not species or arguments.get("dbh_m") is None or target_year is None:
        raise ValueError("predict_species必须提供species、dbh_m和target_year")
    base_year = int(arguments.get("base_year", 2021))
    _, metadata = load_bundle()
    values = dict(metadata["numeric_feature_medians"])
    values.update({key: arguments[key] for key in values if arguments.get(key) is not None})
    values.update({"species": species, "dbh_m": float(arguments["dbh_m"])})
    feature_row = pd.DataFrame([values], columns=FEATURE_COLUMNS)
    result = _prediction_payload(feature_row, base_year, [2025, int(target_year)])
    database = _prepared_database(base_year)
    scenario = pd.Series({"tree_id": "scenario-001", "species": species, "plot_id": ""})
    record = _output_record(scenario, result, int(target_year), _shape_relations(database))
    return _artifact_result("predict_species", [record], result, int(target_year), species=species)


def _predict_species_trees(arguments: dict) -> dict:
    species = str(arguments.get("species") or "").strip()
    if not species or arguments.get("target_year") is None:
        raise ValueError("predict_species_trees必须提供species和target_year")
    target_year = int(arguments["target_year"])
    base_year = int(arguments.get("base_year", 2021))
    frame = _prepared_database(base_year)
    selected = frame.loc[frame["species"].astype(str).eq(species)]
    if selected.empty:
        raise ValueError("数据库中没有可预测的该树种单木")
    relations = _shape_relations(frame)
    records = []
    last_result = None
    for _, tree in selected.iterrows():
        last_result = _prediction_payload(
            pd.DataFrame([tree[FEATURE_COLUMNS]]), base_year, [2025, target_year]
        )
        records.append(_output_record(tree, last_result, target_year, relations))
    assert last_result is not None
    return _artifact_result(
        "predict_species_trees", records, last_result, target_year, species=species,
    )


def _predict_sample(arguments: dict) -> dict:
    target_year = arguments.get("target_year")
    if target_year is None:
        raise ValueError("predict_sample必须提供target_year")
    target_year = int(target_year)
    base_year = int(arguments.get("base_year", 2021))
    sample_size = int(arguments.get("sample_size", 50))
    if sample_size < 1:
        raise ValueError("sample_size必须至少为1")
    seed = int(arguments.get("seed", target_year))
    frame = _prepared_database(base_year)
    if sample_size > len(frame):
        raise ValueError("sample_size超过数据库可预测树木数量")
    sampled = frame.sample(n=sample_size, random_state=seed)
    relations = _shape_relations(frame)
    records = []
    last_result = None
    for _, tree in sampled.iterrows():
        last_result = _prediction_payload(
            pd.DataFrame([tree[FEATURE_COLUMNS]]), base_year, [2025, target_year]
        )
        records.append(_output_record(tree, last_result, target_year, relations))
    assert last_result is not None
    return _artifact_result(
        "predict_sample", records, last_result, target_year,
        sample_size=sample_size, seed=seed,
    )


def predict(arguments: dict) -> dict:
    try:
        operation = arguments["operation"]
        if operation == "status":
            _, metadata = load_bundle()
            return {
                "operation": "status",
                "ready": True,
                "model_name": metadata["model_name"],
                "model_version": metadata["model_version"],
                "training_rows": metadata["training_rows"],
                "species_count": metadata["species_count"],
                "training_years": metadata["training_years"],
                "validation": metadata["cross_validation"]["overall"],
                "experimental": True,
                "output_format": "xlsx",
            }
        if operation == "predict_tree":
            return _predict_tree(arguments)
        if operation == "predict_species":
            return _predict_species(arguments)
        if operation == "predict_species_trees":
            return _predict_species_trees(arguments)
        if operation == "predict_sample":
            return _predict_sample(arguments)
        raise ValueError("不支持的预测操作")
    except (FileNotFoundError, KeyError, RuntimeError, ValueError, psycopg.Error) as exc:
        return {"isError": True, "error": str(exc), "artifacts": []}
