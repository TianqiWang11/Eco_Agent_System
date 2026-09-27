import base64
import json
from io import BytesIO

import pandas as pd
import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from src.agent.tools.local_tools.predict import PREDICT_SCHEMA, predict
from src.agent.tools.local_tools.predict.model_store import METADATA_PATH, MODEL_PATH
from src.agent.tools.local_tools.predict.tool import (
    OUTPUT_COLUMN_WIDTHS,
    _artifact_result,
    _excel_artifact,
    output_headers,
)
from src.agent.tools.local_tools.predict.training import (
    CATEGORICAL_FEATURES,
    FEATURE_COLUMNS,
    GROUP_COLUMN,
)
from src.agent.tools.local_tools.predict import tool as predict_tool


def test_species_is_the_only_categorical_model_feature():
    assert CATEGORICAL_FEATURES == ["species"]
    assert "species" in FEATURE_COLUMNS
    assert "plot_id" not in FEATURE_COLUMNS
    assert GROUP_COLUMN == "plot_id"


def test_agent_prediction_schema_does_not_accept_plot_id_or_training():
    assert "plot_id" not in PREDICT_SCHEMA["properties"]
    assert "train" not in PREDICT_SCHEMA["properties"]["operation"]["enum"]
    assert set(PREDICT_SCHEMA["properties"]["operation"]["enum"]) == {
        "status", "predict_tree", "predict_species", "predict_species_trees", "predict_sample"
    }
    assert "target_year" in PREDICT_SCHEMA["properties"]
    assert "target_years" not in PREDICT_SCHEMA["properties"]


def test_trained_model_bundle_loads_without_writing_output():
    assert MODEL_PATH.is_file() and METADATA_PATH.is_file()
    metadata = json.loads(METADATA_PATH.read_text(encoding="utf-8"))

    status = predict({"operation": "status"})

    assert status["ready"] is True
    assert status["model_version"] == metadata["model_version"]
    assert status["output_format"] == "xlsx"
    assert not any(MODEL_PATH.parent.glob("*.xlsx"))
    assert not any(MODEL_PATH.parent.glob("*.csv"))


def test_prediction_excel_is_generated_from_the_code_template():
    headers = output_headers(2030)
    assert headers == [
        "编号",
        "物种名称",
        "预测年份",
        "样地编号",
        "predicted_dbh_m_2030",
        "estimated_tree_height_m_2030",
        "estimated_crown_diameter_m_2030",
        "estimated_crown_diameter_ns_m_2030",
        "estimated_crown_diameter_ew_m_2030",
        "estimated_crown_area_m2_2030",
        "estimated_crown_volume_m3_2030",
    ]
    record = dict(zip(headers, ["000001", "木荷", 2030, "P001", *([1.25] * 7)]))
    artifact = _excel_artifact([record], 2030, "dbh_predictions_2030_random_1.xlsx")
    workbook = load_workbook(BytesIO(base64.b64decode(artifact["content_base64"])))
    sheet = workbook.active
    assert sheet is not None

    assert workbook.sheetnames == ["predictions_2030"]
    assert [cell.value for cell in sheet[1]] == headers
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:K2"
    assert sheet["A2"].value == "000001"
    assert sheet["A2"].number_format == "@"
    assert sheet["E2"].number_format == "0.0000"
    assert [sheet.column_dimensions[get_column_letter(index)].width for index in range(1, 12)] == OUTPUT_COLUMN_WIDTHS


def test_random_sample_artifact_uses_the_final_filename_pattern():
    headers = output_headers(2030)
    record = dict(zip(headers, ["000001", "木荷", 2030, "P001", *([1.25] * 7)]))
    result = {"model_version": "test"}

    payload = _artifact_result("predict_sample", [record], result, 2030, seed=2030)

    assert payload["artifacts"][0]["filename"] == "dbh_predictions_2030_random_1.xlsx"
    action = payload["ue_actions"][0]
    assert action["type"] == "apply_prediction"
    assert action["target_id"] == "000001"
    assert action["dbh_m"] == 1.25
    assert action["focus"] is True


def test_species_prediction_requires_explicit_year():
    with pytest.raises(ValueError, match="target_year"):
        predict_tool._predict_species_trees({"species": "木荷"})


def test_species_prediction_uses_every_matching_tree_and_requested_year(monkeypatch):
    frame = pd.DataFrame([
        {"tree_id": f"{index:06d}", "species": "木荷", **{key: 1 for key in FEATURE_COLUMNS}}
        for index in range(51)
    ] + [{"tree_id": "other", "species": "其他", **{key: 1 for key in FEATURE_COLUMNS}}])
    frame["species"] = ["木荷"] * 51 + ["其他"]
    captured = []
    monkeypatch.setattr(predict_tool, "_prepared_database", lambda base_year: frame)
    monkeypatch.setattr(predict_tool, "_shape_relations", lambda database: {})

    def fake_prediction(features, base_year, years):
        captured.append((base_year, years))
        return {"model_version": "test"}

    monkeypatch.setattr(predict_tool, "_prediction_payload", fake_prediction)
    monkeypatch.setattr(
        predict_tool, "_output_record",
        lambda tree, result, year, relations: {"tree_id": tree["tree_id"], "year": year},
    )
    monkeypatch.setattr(
        predict_tool, "_artifact_result",
        lambda operation, records, result, year, **extra: {
            "operation": operation, "records": records, "year": year, **extra,
        },
    )

    result = predict_tool._predict_species_trees({"species": "木荷", "target_year": 2042})

    assert result["operation"] == "predict_species_trees"
    assert result["year"] == 2042
    assert len(result["records"]) == 51
    assert result["records"][-1]["tree_id"] == "000050"
    assert captured == [(2021, [2025, 2042])] * 51
